"""Package generated PNGs without modifying their pixels.

Run: uv run --isolated --no-project --with python-pptx --with reportlab python package_slides.py
"""

from datetime import datetime, timezone
from hashlib import sha256
from html import escape
import json
from pathlib import Path
from shutil import copy2
import struct
from zipfile import ZipFile, ZIP_DEFLATED

from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parent
SPECS = json.loads((ROOT / "slides.json").read_text())
SOURCE = ROOT.parent / "genspark_steel_sky"

NOTES = {
    1: "고정 카메라와 컴퓨터 비전, 역기구학, FSM을 연결한 블록 조작 시스템입니다. 실제 장비에서 발생한 인식과 접촉 문제에 어떻게 대응했는지, 그리고 실행 기록으로 무엇을 확인했는지 소개하겠습니다.",
    2: "과제는 같은 규격의 블록 다섯 개를 조작하는 두 미션입니다. 1차는 180초 안에 20×10cm 지정 영역에 배치하고, 2차는 300초 안에 적재한 뒤 5초 유지해야 합니다. 1차는 경계 2cm 걸침, 세워짐, 일부 겹침을 허용하지만 적재는 인정하지 않습니다. 사진은 작업 영역 등록 장면입니다. 현재 이동 제어를 실기 실행했고, 적재 실기는 남아 있습니다.",
    3: "이 사진은 9월 2일 시연 영상에서 가져온 순차 동작 사례입니다. 카메라에서 블록을 검출하고 파지, 운반, 배치로 이어지는 동작을 실제 장비에서 연결했습니다. 이 장은 동작 연결을 보여주는 사례이며, 영상 길이나 프레임의 재생 시각을 공식 미션 완료시간으로 해석하지 않습니다. 선택적으로 한 사이클의 영상을 20~30초 별도로 재생할 수 있습니다.",
    4: "초기에는 ACT와 SmolVLA를 시도했습니다. ACT 단일 블록 이동은 50회 중 29회 성공했으며, SmolVLA는 목표 근처 접근을 관찰했지만 안정적인 파지는 확보하지 못했습니다. 실제 실패에서 관측 입력, 동작 계획, 접촉 중 어느 단계가 문제인지 나누어 확인할 필요가 있었습니다. 고정된 평면과 정해진 블록이라는 조건을 이용해 최종적으로 CV, IK, FSM을 연결했습니다. 동일 조건 비교를 하지 않았으므로 성공률이 개선됐다는 결론으로 제시하지 않습니다.",
    5: "상위 흐름은 카메라, 블록 검출, 베이스 좌표 변환과 IK, FSM, 로봇입니다. 웹 뷰어는 영상과 측정 상태를 관찰하는 별도 가지이며 화면 오버레이가 제어 명령으로 돌아가는 구조는 아닙니다. FSM은 선택, 파지, 확인, 운반, 배치를 반복합니다. 핵심은 VERIFY에서 파지가 확인된 경우에만 운반으로 진행하는 조건입니다. 미확인 시에는 후보 재시도나 재선택으로 돌아갑니다. 적재는 PLACE 전략을 바꾸는 구조이며 실기 검증은 별도로 남아 있습니다.",
    6: "검출은 색 범위로 후보를 만든 다음 형상과 영역 조건을 함께 적용합니다. 빨간 블록과 경계 테이프는 색이 비슷해 색만으로는 충분하지 않습니다. 면적, 종횡비, 채움 정도로 형상을 검사하고, 지정 영역 안이나 작업 범위 밖 후보를 제외합니다. 사진은 작업 영역 등록 화면이며 필터의 적용 전후 비교 결과가 아닙니다. 웹 뷰어는 검출 결과, 후보, FK 위치를 관찰하는 데 사용했습니다. 검출 정확도나 속도 개선율은 따로 측정하지 않았습니다.",
    7: "좌표 수집은 같은 대응점에서 두 단계를 수행합니다. 블록 윗면에 기준점을 맞추고 관절값과 FK를 기록한 뒤, 팔을 치워 영상에서 블록 중심을 확인합니다. 이를 통해 픽셀과 로봇 베이스의 XY 좌표를 대응시킵니다. 9월 8일 재수집한 9점의 적합 RMS는 8.18mm, 최대 LOO는 26.64mm로 개발 목표에 미달했습니다. FK는 관절값으로 계산한 모델 기준점이므로 이 숫자를 실제 그리퍼 턱의 독립적인 위치 정밀도로 해석하지 않습니다. 이러한 잔여 오차를 전제로 파지 확인과 재시도를 연결했습니다.",
    8: "좌표가 주어져도 모든 자세가 실행 가능한 것은 아니므로 후보별 IK 도달성을 검사합니다. 초기에는 고정된 방향을 유지하려다 손목이 비틀리고 과열 보호 정지가 발생했습니다. 이후 작업 위치에서 중립에 가까운 손목 방향을 선택하고, 블록의 90도 대칭을 이용해 면을 정렬하도록 바꿨습니다. 여기 그림은 방향 선택을 설명하는 개념도입니다. 실측 온도 개선량이나 과열 재발 방지 효과를 정량적으로 검증했다는 뜻은 아닙니다.",
    9: "중심점이 실행 불가능하더라도 바로 포기하지 않고 도달성 검사를 통과하는 대체 후보를 검토합니다. 접근과 닫기 이후에는 그리퍼 위치와 부하를 함께 확인합니다. 확인되면 운반하고 미확인이면 다음 후보, 순회 보류 또는 재선택으로 이어집니다. 오른쪽 파지 관측 사례는 위치 20.99와 절대 부하 500, 빈손 사례는 위치 2.49와 부하 44입니다. 위치는 서보 정규화값으로 mm가 아니며 두 사례는 위치와 방향이 다릅니다. 센서 판정 정확도나 같은 조건의 비교 실험으로 해석하지 않습니다. 모델상 목표 거리가 작아도 실제 파지는 실패할 수 있어 도달과 접촉을 구분했습니다.",
    10: "9월 8일 전체 시험은 외부 180초 제한, 30Hz 조건의 한 번 실행입니다. 후보 실행 시작은 14회였고 센서 판정은 HELD 4회, EMPTY 9회, 판정 전 중단 1회였습니다. 슬롯 해제 동작은 네 번이지만 마지막 블록이 남아 전체 미션은 미완료입니다. 센서 판정, 슬롯 해제 횟수, 규정에 따른 인정 블록 수는 서로 다른 항목입니다. 인정 블록 수를 집계하지 않았으므로 4개 성공이나 성공률 80%로 바꿀 수 없습니다. 외부 제한에는 초기화가 포함되며 정확한 물리적 정지 시각은 기록하지 않았습니다.",
    11: "같은 실행의 사건 순서를 보면 재시도가 파지 회복에 기여한 사례와 시간 비용을 함께 볼 수 있습니다. 한 블록은 첫 PICK에 51.5초가 걸렸고, 재선택 뒤 파지를 확인해 최초 선택부터 해제까지 69.8초가 걸렸습니다. 평균 블록당 예산인 36초를 초과합니다. 36초는 180초를 다섯 블록으로 나눈 값이며 실제 강제 중단 설정값이 아닙니다. 그림의 사건 간 거리는 시간에 비례하지 않습니다. 재시도 적용 전후 비교는 하지 않았으므로 평균 성공률 개선으로 일반화하지 않습니다. 다음에는 후보별 시간 비용과 실패 사유를 함께 관리해야 합니다.",
    12: "이번 과제에서 인식, 좌표, IK, FSM을 실제 장비에 통합하고 파지 확인과 실패 대응, 관제 도구, 실행 기록을 연결했습니다. 이동 미션은 같은 조건의 반복 평가와 재시도 시간 관리가 남아 있습니다. 적재는 접근과 하강 자세 등록, 통합 실행, 5초 유지 검증이 필요합니다. 다음 단계로 관측과 행동, 검증된 결과를 연결한 학습용 실행 기록으로 확장할 수 있습니다. 자동 초기화나 무인 데이터 수집은 아직 완성한 성과가 아닙니다.",
    13: "질의응답용 조건표입니다. 초기 5점과 확장 9점 RMS는 사용자가 제공한 실측값입니다. 15점과 재수집 9점을 포함해 서로 다른 수집 조건이므로 점 수 증가의 인과적 개선 추세로 해석하지 않습니다. 전체 시험은 30Hz이고 이후 단일 블록 시험은 45Hz로 전체 미션 재평가는 없습니다. 카메라 재검사는 약 60.5초의 표본별 p95 최댓값 0.514px이며 10분 검증 결과가 아닙니다. 실험 ID는 task1_20260908_030956입니다.",
    14: "질의응답용 기여 구분입니다. OpenCV의 영상 연산, PlaCo의 FK/IK, LeRobot의 장비 통신을 활용했습니다. 팀은 과제 조건에 맞춘 검출과 영역 필터, 좌표 수집 도구, 자세와 후보 처리, 파지 확인과 FSM, 웹 시각화, 실행 환경과 평가 기록을 통합했습니다. 역할은 사용자가 확정한 최신 분담을 따릅니다. 장비 통합, 데이터 수집과 실패 관찰, 접근법 논의와 결과 검토는 공동 작업입니다.",
}


def main():
    missing = [s["output"] for s in SPECS if not Path(s["output"]).is_file()]
    if missing:
        raise SystemExit("Missing slides: " + ", ".join(missing))

    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    pdf = Canvas(str(ROOT / "steel_sky_14.pdf"), pagesize=(1152, 648))
    pdf.setTitle("Steel / Sky - Team Sim2Real, 14 slides")
    pdf.setAuthor("Team Sim2Real")
    manifest = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "generator": "built-in image_gen",
        "format": "14 raster-image slides; 12 main + 2 appendix",
        "main_talk_seconds": 560,
        "transition_buffer_seconds": 40,
        "palette": ["#355872", "#7AAACE", "#9CD5FF", "#F7F8F0"],
        "images_modified_by_packaging": False,
        "slides": [],
    }
    notes = ["# 발표자 노트\n", "본문 9분 20초 + 전환 여유 40초. 시연과 질의응답은 별도.\n"]
    thumbs = []
    for spec in SPECS:
        number = spec["number"]
        file = Path(spec["output"])
        data = file.read_bytes()
        width, height = struct.unpack(">II", data[16:24])
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(str(file), 0, 0, width=prs.slide_width, height=prs.slide_height)
        picture.name = f"{number:02d} {spec['slug']} — imagegen raster slide"
        timing = f"목표 {spec['seconds']}초" if spec["seconds"] else "질의응답용 부록"
        slide.notes_slide.notes_text_frame.text = f"{number:02d} / {timing}\n\n{NOTES[number]}"
        pdf.drawImage(str(file), 0, 0, width=1152, height=648, preserveAspectRatio=False)
        pdf.showPage()
        notes.append(f"## {number:02d}. {spec['slug']} / {timing}\n\n{NOTES[number]}\n")
        relative = file.relative_to(ROOT).as_posix()
        metadata = ROOT / "metadata" / f"{number:02d}.json"
        manifest["slides"].append({
            "number": number, "file": relative, "width": width, "height": height,
            "sha256": sha256(data).hexdigest(), "seconds": spec["seconds"],
            "prompt": f"prompts/{number:02d}_{spec['slug']}.txt",
            "metadata": metadata.relative_to(ROOT).as_posix() if metadata.exists() else None,
        })
        thumbs.append(f'<a href="{escape(relative)}" target="_blank"><img src="{escape(relative)}" alt="{number:02d} {escape(spec["slug"])}" loading="lazy"><span>{number:02d} / {escape(timing)}</span></a>')

    prs.save(ROOT / "steel_sky_14.pptx")
    pdf.save()
    (ROOT / "speaker_notes.md").write_text("\n".join(notes))
    (ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    (ROOT / "preview.html").write_text('''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Steel / Sky · 14 slides</title><style>body{background:#F7F8F0;color:#355872;font-family:system-ui,sans-serif;margin:32px}h1{font-size:32px}main{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:22px}a{color:inherit;text-decoration:none}img{width:100%;display:block;border:1px solid #94aebc}span{display:block;margin:8px 0 4px;font-size:14px}@media(max-width:1100px){main{grid-template-columns:repeat(2,minmax(0,1fr))}}@media(max-width:650px){main{grid-template-columns:1fr}}</style><h1>Steel / Sky</h1><p>본문 12장 + 부록 2장 / 이미지를 클릭하면 원본 크기로 열립니다.</p><main>''' + "".join(thumbs) + "</main></html>")
    originals = ROOT / "source_photos"
    originals.mkdir(exist_ok=True)
    photo_sources = {}
    for spec in SPECS:
        for source_path in spec["referenced_image_paths"][1:]:
            src = Path(source_path)
            dest = originals / src.name
            if not dest.exists():
                copy2(src, dest)
            photo_sources[src.name] = {"source": str(src), "sha256": sha256(src.read_bytes()).hexdigest()}
    (originals / "manifest.json").write_text(json.dumps(photo_sources, ensure_ascii=False, indent=2) + "\n")
    (ROOT / "README.md").write_text("""# Steel / Sky 발표 이미지 14장

- `slides/`: 개별 PNG 14장, 본문 01~12 / 부록 13~14.
- `steel_sky_14.pptx`: 16:9 이미지형 PPTX, 각 장의 발표자 노트 포함.
- `steel_sky_14.pdf`: 동일한 14장 PDF.
- `preview.html`: 브라우저에서 전체 슬라이드를 보고 원본을 여는 목록.
- `speaker_notes.md`: 발표자 노트와 시간 배분.
- `prompts/`, `slides.json`, `metadata/`: 생성/수정 요청문과 검토 기록.
- `source_photos/`: 참고에 사용한 원본 사진 사본과 해시.

제작 방식: built-in image_gen으로 각 슬라이드를 생성하고 필요한 부분을 같은 도구로 수정함. 생성된 PNG의 픽셀을 후처리하지 않고 PPTX/PDF에 배치함. 이미지 비율은 생성기 출력에 따라 16:9에 근접하며 PPTX/PDF 페이지는 정확히 16:9임.

이미지형 PPTX이므로 슬라이드 안의 글자/도형은 개별 편집 요소가 아님. 발표자 노트는 편집 가능함. 편집 가능한 PPT 제작은 기존 Genspark 전달 지침을 사용하면 됨.

사진은 제공된 실제 장면을 참고했지만 이미지 생성 과정에서 다시 표현된 디자인 시안이므로, 원본 픽셀과 동일한 실험 증거로 취급하지 않음. 실제 경계/접촉/좌표를 판정할 때는 source_photos와 최종보고서 원본을 사용함.

성능 표기: ACT 29/50은 초기 단일 블록 결과. 최종 전체 시험의 해제 4회는 인정 블록 수나 성공률 80%를 뜻하지 않음. 전체 미션 미완료, 적재 실기 미수행을 유지함.

PNG 묶음 ZIP은 이미지/프롬프트/노트/참고 사진을 포함하고, 중복 용량을 줄이기 위해 별도로 제공하는 PPTX/PDF는 포함하지 않음.
""")
    archive = ROOT.parent / (ROOT.name + "_png.zip")
    with ZipFile(archive, "w", ZIP_DEFLATED) as zipped:
        for file in sorted(ROOT.rglob("*")):
            if file.is_file() and file.suffix not in {".pptx", ".pdf"}:
                zipped.write(file, Path(ROOT.name) / file.relative_to(ROOT))
    with ZipFile(archive) as zipped:
        assert zipped.testzip() is None
    reopened = Presentation(ROOT / "steel_sky_14.pptx")
    assert len(reopened.slides) == 14
    assert all(len(slide.shapes) == 1 for slide in reopened.slides)
    assert all(slide.notes_slide.notes_text_frame.text for slide in reopened.slides)
    print(json.dumps({"slides": 14, "main_talk_seconds": 560, "files": {
        file.name: file.stat().st_size for file in [ROOT / "steel_sky_14.pptx", ROOT / "steel_sky_14.pdf", archive]
    }}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
